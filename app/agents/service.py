from __future__ import annotations

import io
import json
import re
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.crud import get_ai_model, get_existing_conversation_for_model
from app.models import AIAgent, AIAgentFile, AIAgentRun, AIAgentStep
from app.providers import send_messages_with_model
from app.schemas import ChatMessage, CostUsage, TokenUsage
from app.settings import get_settings

from .schemas import AgentRunRequest, AgentRunResponse


TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".php", ".java", ".cs",
    ".go", ".rs", ".rb", ".sql", ".html", ".css", ".scss", ".vue",
    ".svelte", ".sh", ".ps1", ".toml", ".ini", ".conf",
}
BLOCKED_FILENAMES = {".env", ".env.local", ".env.production", "credentials.json"}
MAX_AGENT_FILE_MB = 12
MAX_EXTRACTED_CHARS = 150_000


def get_agent(db: Session, slug: str) -> AIAgent:
    agent = db.execute(select(AIAgent).where(
        AIAgent.slug == slug,
        AIAgent.is_active.is_(True),
        AIAgent.deleted_at.is_(None),
    )).scalar_one_or_none()
    if not agent:
        raise ValueError("Agent not found or disabled")
    return agent


def serialize_agent(agent: AIAgent) -> dict[str, Any]:
    return {
        "id": agent.id, "slug": agent.slug, "name_ar": agent.name_ar,
        "name_en": agent.name_en, "description": agent.description,
        "primary_model_tool": agent.primary_model_tool,
        "allowed_tools": agent.allowed_tools or [],
        "input_schema": agent.input_schema or {},
        "output_schema": agent.output_schema or {},
        "default_parameters": agent.default_parameters or {},
        "limits": {
            "max_steps": agent.max_steps, "max_files": agent.max_files,
            "max_execution_seconds": agent.max_execution_seconds,
            "default_budget_usd": float(agent.default_budget_usd)
            if agent.default_budget_usd is not None else None,
        },
        "is_recommended": bool(agent.is_recommended),
        "sort_order": agent.sort_order,
        "endpoint": f"/tasks/agents/{agent.slug}/run",
    }


def _safe_filename(filename: str) -> str:
    base = Path(str(filename or "file")).name
    suffix = Path(base).suffix.lower()
    stem = Path(base).stem
    safe_stem = re.sub(r"[^A-Za-z0-9\u0600-\u06FF_-]+", "_", stem).strip("._-")
    safe_suffix = re.sub(r"[^A-Za-z0-9.]", "", suffix)
    return f"{(safe_stem or 'file')[:160]}{safe_suffix[:15]}"


def _looks_like_docx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            return "[Content_Types].xml" in names and "word/document.xml" in names
    except (OSError, zipfile.BadZipFile):
        return False


def _extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if Path(filename).name.lower() in BLOCKED_FILENAMES:
        raise ValueError(f"Sensitive file is not allowed: {filename}")
    is_docx = suffix == ".docx" or _looks_like_docx(data)
    if is_docx:
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("python-docx is required for DOCX agent files") from exc
        try:
            document = Document(io.BytesIO(data))
        except UnicodeDecodeError as exc:
            raise ValueError(
                "DOCX contains invalid encoded XML; open it in Word or LibreOffice and Save As DOCX"
            ) from exc
        except (OSError, ValueError, zipfile.BadZipFile) as exc:
            raise ValueError("Uploaded DOCX file is invalid or unsupported") from exc
        sections: list[str] = []
        paragraph_text = "\n".join(
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        if paragraph_text:
            sections.append(paragraph_text)
        for table_index, table in enumerate(document.tables, start=1):
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                sections.append(f"[Table {table_index}]\n" + "\n".join(rows))
        return "\n\n".join(sections)
    if suffix in TEXT_EXTENSIONS:
        return data.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("pypdf is required for PDF agent files") from exc
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(
            f"[Page {index}]\n{page.extract_text() or ''}"
            for index, page in enumerate(reader.pages, start=1)
        )
    raise ValueError(f"Unsupported agent file type: {suffix or 'unknown'}")


async def _prepare_files(
    db: Session, run: AIAgentRun, files: list[UploadFile], max_files: int
) -> tuple[str, list[dict[str, Any]]]:
    if len(files) > max_files:
        raise ValueError(f"Agent accepts at most {max_files} files")
    storage_dir = Path(get_settings().GENERATED_FILES_DIR) / "agents" / run.run_uuid / "input"
    storage_dir.mkdir(parents=True, exist_ok=True)
    extracted_sections: list[str] = []
    file_items: list[dict[str, Any]] = []
    total_chars = 0
    for position, upload in enumerate(files, start=1):
        filename = _safe_filename(upload.filename or f"file-{position}")
        data = await upload.read()
        if not data:
            raise ValueError(f"Uploaded file is empty: {filename}")
        if len(data) > MAX_AGENT_FILE_MB * 1024 * 1024:
            raise ValueError(f"{filename} exceeds {MAX_AGENT_FILE_MB} MB")
        text = _extract_text(filename, data)
        remaining = MAX_EXTRACTED_CHARS - total_chars
        if remaining <= 0:
            break
        excerpt = text[:remaining]
        total_chars += len(excerpt)
        unique_name = f"{position:02d}-{uuid.uuid4().hex[:8]}-{filename}"
        path = storage_dir / unique_name
        path.write_bytes(data)
        item = {
            "filename": filename, "content_type": upload.content_type,
            "size_bytes": len(data), "extracted_chars": len(excerpt),
        }
        db.add(AIAgentFile(
            run_id=run.id, role="input", filename=filename,
            storage_path=str(path), content_type=upload.content_type,
            size_bytes=len(data), file_metadata={"extracted_chars": len(excerpt)},
            created_at=datetime.utcnow(),
        ))
        file_items.append(item)
        extracted_sections.append(f"\n===== SOURCE: {filename} =====\n{excerpt}")
    db.flush()
    return "\n".join(extracted_sections), file_items


def _effective_parameters(agent: AIAgent, req: AgentRunRequest) -> dict[str, Any]:
    values = {**(agent.default_parameters or {})}
    values.update(req.state.parameters.model_dump(exclude_none=True))
    values["max_steps"] = min(int(values.get("max_steps") or agent.max_steps), agent.max_steps)
    values["max_tokens"] = min(
        int(values.get("max_tokens") or 6000), get_settings().GENERAL_TOOLS_MAX_OUTPUT_TOKENS
    )
    budget = values.get("max_budget_usd")
    if budget is None and agent.default_budget_usd is not None:
        budget = float(agent.default_budget_usd)
    values["max_budget_usd"] = float(budget) if budget is not None else None
    return values


def _task_instructions(agent: AIAgent, parameters: dict[str, Any]) -> str:
    visible = {
        key: value for key, value in parameters.items()
        if key not in {"max_budget_usd", "max_steps", "max_tokens", "temperature", "quality_mode"}
    }
    return (
        "Complete the request using this workflow:\n"
        "1. Identify the objective, constraints, and missing assumptions.\n"
        "2. Analyze the supplied file excerpts when present.\n"
        "3. Produce the complete final deliverable in the requested language.\n"
        "4. Do not output internal chain-of-thought, planning notes, or tool-call simulation.\n"
        f"Agent configuration: {json.dumps(visible, ensure_ascii=False)}"
    )


def _actual_cost(raw_usage: dict | None) -> CostUsage | None:
    try:
        value = (raw_usage or {}).get("cost")
        return CostUsage(total_cost=float(value), currency="USD") if value is not None else None
    except (TypeError, ValueError):
        return None


async def run_agent(
    db: Session, agent: AIAgent, req: AgentRunRequest,
    files: list[UploadFile], request_id: str,
) -> AgentRunResponse:
    get_existing_conversation_for_model(
        db, req.user_id, req.model_id, req.conversation_uuid
    )
    model = get_ai_model(
        db, req.selected_model_id,
        tool_key=agent.primary_model_tool, operation="text_generation",
    )
    if model.provider.lower() != "openrouter":
        raise ValueError("Agent primary model must use OpenRouter")
    allowed_models = {int(value) for value in (agent.allowed_model_ids or [])}
    if allowed_models and int(model.id) not in allowed_models:
        raise ValueError("Selected model is not allowed for this agent")
    if agent.slug == "document-knowledge" and not files:
        raise ValueError("At least one document is required for document-knowledge")

    now = datetime.utcnow()
    run = AIAgentRun(
        run_uuid=str(uuid.uuid4()), agent_id=agent.id, user_id=req.user_id,
        model_id=req.model_id, selected_model_id=req.selected_model_id,
        conversation_uuid=req.conversation_uuid, user_message=req.user_message,
        status="running", input_state=req.state.model_dump(), total_steps=2,
        completed_steps=0, started_at=now, created_at=now, updated_at=now,
    )
    db.add(run)
    db.flush()
    prepare_step = AIAgentStep(
        run_id=run.id, step_number=1, step_key="prepare_context",
        title="Validate request and prepare files", tool_key="document_extract" if files else None,
        operation="extract" if files else "validate", status="running",
        tool_input={"file_count": len(files)}, started_at=now, created_at=now,
    )
    db.add(prepare_step)
    db.flush()
    db.commit()

    try:
        parameters = _effective_parameters(agent, req)
        extracted_context, file_items = await _prepare_files(db, run, files, agent.max_files)
        prepare_step.status = "completed"
        prepare_step.tool_output = {"files": file_items, "context_chars": len(extracted_context)}
        prepare_step.completed_at = datetime.utcnow()
        run.completed_steps = 1

        generate_step = AIAgentStep(
            run_id=run.id, step_number=2, step_key="generate_result",
            title="Generate final deliverable", tool_key=agent.primary_model_tool,
            operation="text_generation", selected_model_id=model.id,
            status="running", tool_input={"parameters": parameters},
            started_at=datetime.utcnow(), created_at=datetime.utcnow(),
        )
        db.add(generate_step)
        db.commit()

        user_content = f"{_task_instructions(agent, parameters)}\n\nUSER REQUEST:\n{req.user_message}"
        if extracted_context:
            user_content += f"\n\nSUPPLIED FILE EXTRACTS:\n{extracted_context}"
        reasoning = {"fast": None, "balanced": "low", "high": "medium"}.get(
            str(parameters.get("quality_mode") or "balanced"), "low"
        )
        result = await send_messages_with_model(
            model_key="writer_pro", model_override=model.provider_model_id,
            messages=[
                ChatMessage(role="system", content=agent.system_prompt),
                ChatMessage(role="user", content=user_content),
            ],
            temperature_override=float(
                0.4 if parameters.get("temperature") is None else parameters["temperature"]
            ),
            max_tokens_override=int(parameters["max_tokens"]),
            exclude_reasoning=True, reasoning_effort=reasoning,
        )
        usage = TokenUsage(
            input_tokens=result.input_tokens, output_tokens=result.output_tokens,
            total_tokens=result.total_tokens,
        )
        cost = _actual_cost(result.raw_usage)
        budget = parameters.get("max_budget_usd")

        generate_step.status = "completed"
        generate_step.tool_output = {"content_chars": len(result.content)}
        generate_step.usage_data = usage.model_dump()
        generate_step.cost_data = cost.model_dump() if cost else None
        generate_step.completed_at = datetime.utcnow()
        run.status = "completed"
        run.completed_steps = 2
        run.final_response = result.content
        run.usage_data = usage.model_dump()
        run.cost_data = cost.model_dump() if cost else None
        run.completed_at = datetime.utcnow()
        run.updated_at = datetime.utcnow()
        db.commit()

        return AgentRunResponse(
            agent=agent.slug, run_uuid=run.run_uuid, status=run.status,
            provider="openrouter", model=result.returned_model or model.provider_model_id,
            model_tier=model.tier, user_id=req.user_id, model_id=req.model_id,
            selected_model_id=req.selected_model_id,
            conversation_uuid=req.conversation_uuid, content=result.content,
            usage=usage, cost=cost,
            steps=[
                {"step_number": 1, "key": "prepare_context", "status": "completed"},
                {"step_number": 2, "key": "generate_result", "status": "completed"},
            ],
            files=file_items, request_id=request_id,
            metadata={
                "allowed_tools": agent.allowed_tools or [],
                "effective_parameters": parameters,
                "budget_exceeded": bool(cost and budget and cost.total_cost > budget),
                **result.trace_metadata(),
            },
        )
    except Exception as exc:
        db.rollback()
        persisted = db.execute(select(AIAgentRun).where(AIAgentRun.id == run.id)).scalar_one_or_none()
        if persisted:
            persisted.status = "failed"
            persisted.error_message = str(exc)[:4000]
            persisted.completed_at = datetime.utcnow()
            persisted.updated_at = datetime.utcnow()
            running_steps = db.execute(select(AIAgentStep).where(
                AIAgentStep.run_id == persisted.id, AIAgentStep.status == "running"
            )).scalars().all()
            for step in running_steps:
                step.status = "failed"
                step.error_message = str(exc)[:4000]
                step.completed_at = datetime.utcnow()
            db.commit()
        raise
